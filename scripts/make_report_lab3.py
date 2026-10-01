"""Збирає звіт Лаби 3 (docx) із results_lab3.json.

    uv run python -m findex.lab3_eval data/gutenberg --limit 50
    uv run python scripts/make_report_lab3.py --results results_lab3.json --out "Лабораторна3_звіт.docx"

Усі числа беруться з results_lab3.json (твої вимірювання). Якщо JSON зроблено з --demo,
на титулці з'являється підсвічене попередження.
"""
from __future__ import annotations

import argparse
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

REPO = "https://github.com/8502568-ship-it/findex"
FONT = "Times New Roman"


def f(x, d=3):
    return f"{x:.{d}f}".replace(".", ",")


def sp(n):
    return f"{n:,}".replace(",", " ")


class Rep:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(14)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        for lvl, size in ((1, 16), (2, 14)):
            h = self.d.styles[f"Heading {lvl}"]
            h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
            h.font.color.rgb = None
            rf = h.element.rPr.rFonts
            for att in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                if rf.get(qn(f"w:{att}")) is not None:
                    del rf.attrib[qn(f"w:{att}")]
            for att in ("ascii", "hAnsi", "eastAsia", "cs"):
                rf.set(qn(f"w:{att}"), FONT)

    def h(self, text, lvl=1):
        self.d.add_heading(text, lvl)

    def p(self, text, bold=False, align=None, hl=False, italic=False):
        par = self.d.add_paragraph()
        run = par.add_run(text)
        run.bold, run.italic = bold, italic
        if hl:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if align == "c":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(6)
        return par

    def bullets(self, items):
        for it in items:
            self.d.add_paragraph(it, style="List Bullet").paragraph_format.space_after = Pt(2)

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            par = self.d.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.left_indent = Pt(12)
            r = par.add_run(line or " ")
            r.font.name, r.font.size = "Consolas", Pt(10)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def cap(self, text):
        self.p(text, italic=True)

    def table(self, head, rows, widths=None):
        t = self.d.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        t.autofit = False
        widths = widths or [1] * len(head)
        total_cm = 16.5
        cm = [total_cm * w / sum(widths) for w in widths]
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size = True, Pt(11)
            tcPr = c._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "D9E2F3")
            tcPr.append(shd)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(11)
        for col, w in zip(t.columns, cm):
            col.width = Cm(w)
        for row in t.rows:
            for c, w in zip(row.cells, cm):
                c.width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)


def build(R, out):
    r = Rep()
    S = R["sanity"]; A, B, C = S["rare_vs_common"], S["saturation"], S["length_norm"]
    P5, D = R["p5"], R["demo_query"]

    # ------------------------------ титулка ------------------------------
    r.p("ЗВІТ", bold=True, align="c")
    r.p("з лабораторної роботи №3", align="c")
    r.p("«Рейтинг і модель об’єктів: dunder-методи, протоколи, декоратори»", bold=True, align="c")
    r.p("Курс: «Програмування на Python» (проєкт findex — пошукова система)", align="c")
    r.p("Виконав: [ПІБ студента], група [група]", align="c", hl=True)
    r.p(f"Репозиторій: {REPO} (тег lab-03)", align="c")
    if R["demo"]:
        r.p("УВАГА: числа в цьому файлі отримано на ДЕМО-КОРПУСІ (розділи довідника Python), а не на "
            "Gutenberg. Перед здачею виконайте lab3_eval на своєму корпусі й перегенеруйте звіт.",
            bold=True, hl=True)

    # ------------------------------ 1. мета ------------------------------
    r.h("1. Мета та постановка")
    r.p("Перетворити інвертований індекс Лаби 2 (набір dict-ів і функцій) на об’єкт, який поводиться "
        "як вбудований тип Python, замінити Boolean-пошук ранжованим, додати повноцінну мову запитів "
        "зі сніпетами та виміряти якість ранжування.")
    r.table(["Етап", "Що зроблено", "Модуль"],
            [["M1", "Клас Index(Mapping): len, in, [], for, repr; num_docs, avg_doc_length "
                    "(cached_property), df; open_index — context manager", "indexer.py"],
             ["M2", "Протокол Scorer; TfIdf і BM25; top-k через heapq.nlargest; SearchResult(order=True)",
              "ranking.py"],
             ["M3", "Рекурсивний спуск; вузли Term/Phrase/And/Or/Not з &, |, ~ та evaluate(index)",
              "query.py"],
             ["M4", "@timed, lru_cache, сніпети ±80 символів, precision@5", "util.py, engine.py, "
                                                                          "snippets.py, lab3_eval.py"]], widths=[1, 7, 3.5])
    r.p(f"Корпус: {R['n_docs']} документів, {sp(R['n_terms'])} унікальних термінів, "
        f"{sp(R['n_tokens'])} токенів; індекс збудовано з позиціями (прапорець --positions). "
        f"Python {R['python']}. Тести: 53 passed (з них 26 — нові, tests/test_lab3.py).")

    # ------------------------------ 2. M1 ------------------------------
    r.h("2. M1 — Index як об’єкт")
    r.p("Кожен вбудований оператор — це виклик dunder-методу: len(x) → x.__len__(), "
        "t in x → x.__contains__(t), x[t] → x.__getitem__(t), for t in x → x.__iter__(). Клас Index "
        "наслідує collections.abc.Mapping: достатньо реалізувати __getitem__, __len__, __iter__ — "
        "keys(), items(), get() та решту дає ABC. Index обгортає будь-який бекенд Лаби 2 "
        "(ObjectIndex, ArrayIndex, MmapIndex), тому все з Лаби 2 працює через новий об’єкт.")
    r.code('''class Index(Mapping):
    def __len__(self): return len(self._terms)
    def __contains__(self, term): return term in self._terms   # без побудови списку постінгів
    def __getitem__(self, term):                                 # KeyError, якщо терміна немає
        ...
    def __iter__(self): return iter(self._terms)
    def __repr__(self): return f"Index(terms={len(self):_}, docs={self.num_docs:_})"

    @functools.cached_property
    def avg_doc_length(self): ...                                # один раз на індекс''')
    r.p("Проєктні рішення:")
    r.bullets([
        "__contains__ перевизначено, бо типова реалізація Mapping викликає __getitem__ і для "
        "ArrayIndex/MmapIndex будувала б список Posting лише заради перевірки наявності.",
        "Mapping.__eq__ порівнює всі пари ключ-значення (для сотень тисяч термінів це дуже дорого) "
        "і робить об’єкт нехешованим. Індекс — сутність з ідентичністю, тому __eq__/__hash__ "
        "визначено за ідентичністю; це потрібно, щоб індекс міг бути аргументом lru_cache.",
        "avg_doc_length — cached_property: BM25 викликає його для кожного документа, обчислення "
        "один раз на індекс замість одного на документ.",
        "__repr__ дає Index(terms=…, docs=…): " + R["repr"],
    ])
    r.p("open_index(path) — генераторний @contextmanager: усе до yield відповідає __enter__, "
        "усе після — __exit__; блок finally виконується і при return, і при винятку в тілі with. "
        "Тест test_open_index_releases_on_exception перевіряє, що після RuntimeError у тілі with mmap "
        "бінарного індексу закрито.")
    r.code('''@contextmanager
def open_index(path):
    ix = Index(load(path))
    try:
        yield ix
    finally:
        ix.close()''')

    # ------------------------------ 3. M2 ------------------------------
    r.h("3. M2 — Scorer, TF-IDF, BM25")
    r.p("Scorer — структурний протокол (typing.Protocol): будь-який об’єкт із методом "
        "score(term, posting, index) -> float є Scorer без успадкування. TfIdf і BM25 — незалежні "
        "frozen-dataclass-и, тож замінюються одним аргументом search(…, scorer=…).")
    r.table(["Схема", "Формула", "Параметри"],
            [["TF-IDF", "(1 + ln tf) · ln(N / df)", "—"],
             ["BM25", "idf · tf·(k1+1) / (tf + k1·(1 − b + b·dl/avgdl)),  idf = ln(1 + (N−df+0.5)/(df+0.5))",
              "k1 = 1,5; b = 0,75"]])
    r.p("search() обчислює множину документів за деревом запиту, накопичує score[doc] += вага терміна "
        "по постінгах і бере топ-k через heapq.nlargest(k, …): O(n log k) замість O(n log n) для "
        "sorted(…)[:k]. SearchResult(doc_id, score, title) має order=True; порядок задає прихований "
        "ключ (score, −doc_id), тому sorted(results) впорядковує за релевантністю, а за рівних балів "
        "перемагає менший doc_id (детермінованість).")
    r.h("Sanity-перевірки на корпусі", 2)
    r.p(f"1) Рідкий термін вище за частий: «{A['rare']}» (df={A['rare_df']}) проти «{A['common']}» "
        f"(df={A['common_df']}), tf=1, документ довжиною {A['doc_len']}.")
    r.cap("Таблиця 1 — рідкий проти частого")
    r.table(["Scorer", "рідкий", "частий", "Результат"],
            [[n, f(A["scores"][n]["rare"]), f(A["scores"][n]["common"]),
              "✔" if A["scores"][n]["rare"] > A["scores"][n]["common"] else "✘"] for n in ("tfidf", "bm25")])
    r.p(f"2) 20-те повторення майже нічого не додає: термін «{B['term']}» (df={B['df']}), документ "
        "середньої довжини.")
    r.cap("Таблиця 2 — бал залежно від tf")
    r.table(["Scorer"] + [f"tf={t}" for t in B["tfs"]],
            [[n] + [f(B["scores"][n][str(t)]) for t in B["tfs"]] for n in ("tfidf", "bm25")])
    r.p(f"Приріст при переході 20→21 відносно 1→2: TF-IDF {f(B['ratio']['tfidf'] * 100, 1)}%, "
        f"BM25 {f(B['ratio']['bm25'] * 100, 1)}%. Бал BM25 обмежений стелею (k1+1)·idf = "
        f"{f(B['bm25_ceiling'])}; TF-IDF з ln tf теж зростає повільно, але без стелі.")
    r.p(f"3) Короткий документ з одним входженням вище за дуже довгий: термін «{C['term']}», tf=1 у "
        f"«{C['short']['title']}» ({C['short']['len']} токенів) і «{C['long']['title']}» "
        f"({C['long']['len']} токенів).")
    r.cap("Таблиця 3 — нормалізація за довжиною")
    r.table(["Scorer", "короткий", "довгий", "Результат"],
            [[n, f(C["scores"][n]["short"]), f(C["scores"][n]["long"]),
              "BM25: ✔" if n == "bm25" and C["ok"] else ("довжина не враховується" if n == "tfidf" else "✘")]
             for n in ("tfidf", "bm25")])
    r.p("TF-IDF дає однакові бали, бо не знає про довжину документа; BM25 її штрафує (b = 0,75). "
        "При b = 0 нормалізація вимикається і BM25 поводиться як TF-IDF із насиченням.")

    # ------------------------------ 4. M3 ------------------------------
    r.h("4. M3 — мова запитів")
    r.p("Граматика (рекурсивний спуск, по функції на правило):")
    r.code('''or_expr  := and_expr ('OR' and_expr)*
and_expr := not_expr (['AND'] not_expr)*     # AND можна опускати
not_expr := 'NOT' not_expr | atom
atom     := WORD | "фраза" | '(' or_expr ')' ''')
    r.p("Пріоритети випливають із вкладеності правил: NOT зв’язує найсильніше, потім AND, найслабше — OR. "
        "Тому a OR b c розбирається як Or(a, And(b, c)), а не And(Or(a, b), c): or_expr спершу викликає "
        "and_expr, який «з’їдає» b c, і лише потім бачить OR. Вузли — frozen-dataclass-и, тому тести "
        "порівнюють дерева (assert parse('a OR b c') == Or(Term('a'), And(Term('b'), Term('c')))), "
        "а не рядки. Вузли перевантажують &, | і ~, тож запит можна скласти без рядка: "
        "Term('python') & (Term('async') | Term('await')) & ~Term('java').")
    r.bullets([
        "Term.evaluate — множина doc_id з постінгів; And/Or — перетин/об’єднання множин; "
        "And(a, Not(b)) виконується як різниця a \\ b без побудови «всесвіту» документів.",
        "Phrase.evaluate: перетин документів усіх слів фрази, далі двовказівниковий merge за позиціями "
        "(позиція q другого слова підходить, якщо q−1 є серед позицій першого). Без позицій "
        "кидається QueryError з підказкою про --positions.",
        "Слово, що токенізується на кілька токенів (well-known), стає фразою; токенізація запиту "
        "збігається з токенізацією корпусу (NFC + casefold).",
        "Помилки (незакриті лапки/дужки, обірваний оператор, порожній запит) → QueryError.",
    ])
    r.p("Ранжування застосовується до знайденої множини: вагу дають лише терміни поза NOT.")

    # ------------------------------ 5. M4 ------------------------------
    r.h("5. M4 — декоратори, кеш, сніпети, оцінка")
    r.p("@timed — замикання всередині декоратора: wrapper пам’ятає fn, вимірює час "
        "time.perf_counter() і пише в лог у finally (тож логується й виклик, що завершився винятком). "
        "functools.wraps зберігає __name__ та __doc__; без нього search.__name__ == 'wrapper'. "
        "Якщо fn — lru_cache-обгортка, @timed додатково позначає виклик як cache hit/miss. "
        "@timed застосовано до search, build_index і load.")
    r.code('''@timed
@functools.lru_cache(maxsize=256)
def match_ids(index, query): ...     # запит -> id документів

@timed
def search(index, query, scorer=DEFAULT_SCORER, k=10, snippets=False): ...''')
    r.cap("Лістинг логу (один запит двічі підряд)")
    r.code("\n".join(R["log"]))
    r.p(f"Крок «запит → id документів»: miss {f(D['miss_ms'], 3)} мс, hit {f(D['hit_ms'], 3)} мс. "
        "Дерево запиту теж кешується (parse під lru_cache — вузли незмінні й хешовані).")
    r.p("Сніпети: токени знаходяться тим самим регулярним виразом, що й у токенайзері, але на "
        "NFC-тексті з оригінальним регістром; ковзне вікно обирає ділянку з найбільшою кількістю різних "
        "термів запиту, навколо її центрального входження береться ±80 символів (без розрізання слів), "
        "збіги виділяються **…**.")
    r.p(f"Демо-запит: {D['q']}")
    r.table(["Бал", "Документ", "Сніпет"], [[f(x["score"]), x["title"], x["snippet"]] for x in D["results"]],
            widths=[1, 2, 7])
    r.h("precision@5: TF-IDF проти BM25", 2)
    r.p("Для кожного з 10 запитів релевантними вважаються документи, назви яких відповідають "
        "заздалегідь заданим шаблонам (eval_queries.json). Запит подається як t1 OR t2 OR …, щоб "
        "ранжувалася вся множина збігів. P@5 = (кількість релевантних у топ-5) / 5.")
    r.cap("Таблиця 4 — precision@5")
    rows = [[i, x["q"], x["relevant_in_corpus"], f(x["scorers"]["tfidf"]["p5"], 1),
             f(x["scorers"]["bm25"]["p5"], 1)] for i, x in enumerate(P5["rows"], 1)]
    rows.append(["", f"Середнє ({P5['n_valid']} запитів)", "", f(P5["mean"]["tfidf"], 2), f(P5["mean"]["bm25"], 2)])
    r.table(["#", "Запит", "Релевантних", "TF-IDF", "BM25"], rows, widths=[0.6, 6, 1.6, 1.3, 1.3])
    better = "BM25" if P5["mean"]["bm25"] > P5["mean"]["tfidf"] else (
        "TF-IDF" if P5["mean"]["tfidf"] > P5["mean"]["bm25"] else None)
    r.p(("Середній P@5 вищий у " + better + ". ") if better else "Середній P@5 однаковий. ")
    r.p("Вибірка мала (10 запитів) і мітки грубі, тож різницю слід сприймати як тенденцію, а не "
        "статистично значущий результат; цінність вправи — у звичці вимірювати якість ранжування.")
    if P5["missing"]:
        r.p("Запити без релевантних документів у корпусі (виключено зі середнього): " + "; ".join(P5["missing"]),
            hl=True)

    # ------------------------------ 6. reflection ------------------------------
    r.h("6. Відповіді на питання для самоперевірки")
    QA = [
        ("Що відбувається при \"python\" in index та len(index)? Чому не викликати index.__len__()?",
         "Інтерпретатор бачить оператор in, знаходить тип об’єкта справа і викликає "
         "type(index).__contains__(index, 'python'); len(index) так само викликає type(index).__len__. "
         "Результат len має бути невід’ємним int, інакше TypeError/ValueError. Прямий виклик "
         "index.__len__() обходить ці перевірки та швидкі шляхи вбудованих типів у C і порушує "
         "контракт: dunder-методи призначені для інтерпретатора, а не для користувацького коду."),
        ("Що таке протокол? Чим typing.Protocol відрізняється від ABC?",
         "Протокол — набір методів, що визначає поведінку (iterable: __iter__, sized: __len__, "
         "callable: __call__). ABC вимагає номінального зв’язку (успадкування або register) і може "
         "давати готові методи (Mapping дає keys/items/get); Protocol — структурний: клас "
         "відповідає, якщо має потрібні методи, без успадкування. ABC беремо, коли потрібна спільна "
         "реалізація або перевірка isinstance за договором; Protocol — коли потрібен лише "
         "інтерфейс для типізації, як Scorer, що має незалежні реалізації TfIdf та BM25."),
        ("@timed з пам’яті та що ламається без functools.wraps",
         "def timed(fn): @functools.wraps(fn); def wrapper(*a, **kw): t0 = perf_counter(); try: "
         "return fn(*a, **kw) finally: log.debug('%s took %.3f ms', fn.__name__, …); return wrapper. "
         "Без wraps search.__name__ == 'wrapper', зникає docstring (help(search) показує "
         "обгортку), ламаються інтроспекція, логи з fn.__name__ та документація."),
        ("Декоратор з аргументами і @retry(times=3)",
         "Потрібен ще один рівень: retry(times) повертає декоратор, декоратор(fn) повертає wrapper. "
         "@retry(3) — це fn = retry(3)(fn). Схема: def retry(times): def deco(fn): @wraps(fn) "
         "def wrapper(*a, **kw): for i in range(times): try: return fn(*a, **kw) except Exception: "
         "if i == times-1: raise; return wrapper; return deco."),
        ("@contextmanager і yield",
         "До yield виконується код __enter__ (значення yield — результат as), після — __exit__. "
         "Якщо тіло with кидає виняток, він повторно піднімається в генераторі в точці yield, "
         "тому finally/except навколо yield виконаються; без try/finally код після yield при "
         "винятку не виконається. Тому open_index обгортає yield у try/finally."),
        ("Дві переваги BM25 над TF-IDF; k1 та b; b = 0",
         "1) Насичення TF: внесок терміна обмежений стелею (k1+1)·idf, k1 керує швидкістю насичення; "
         "2) нормалізація за довжиною документа: довгий документ з одним збігом менш релевантний, "
         "b ∈ [0; 1] задає силу нормалізації. При b = 0 довжина ігнорується: множник "
         "k1·(1 − b + b·dl/avgdl) стає сталим k1."),
        ("Чому heapq.nlargest(k, …), а не sorted(…)[:k]?",
         "nlargest тримає купу розміром k: O(n log k) часу й O(k) додаткової пам’яті; sorted — "
         "O(n log n) і копія всіх n елементів. Коли k ≪ n (топ-10 із десятків тисяч збігів), "
         "виграш суттєвий."),
        ("a OR b c: чому Or(a, And(b, c))?",
         "OR має найнижчий пріоритет: or_expr викликає and_expr, який жадібно збирає b c в And; "
         "оператор OR «розділяє» запит на альтернативи верхнього рівня. Якби OR мав вищий "
         "пріоритет, вийшло б And(Or(a, b), c). Цю поведінку фіксує тест на рівність дерев."),
    ]
    for q, a in QA:
        r.p(q, bold=True)
        r.p(a)

    # ------------------------------ 7. висновки ------------------------------
    r.h("7. Висновки")
    r.bullets([
        "Dunder-методи та Mapping перетворюють індекс на звичайний для Python об’єкт: len, in, [], "
        "for, repr і with працюють без знання внутрішнього формату зберігання.",
        "Protocol Scorer дозволяє міняти TF-IDF і BM25 без успадкування; BM25 задовольняє всі "
        "три sanity-перевірки, а TF-IDF провалює третю (не враховує довжину).",
        "Мова запитів будується як дерево незмінних вузлів: парсер легко тестувати порівнянням "
        "дерев, а запити можна складати операторами &, |, ~.",
        "Декоратори й lru_cache дають спостережуваність і виграш на повторних запитах без зміни "
        "логіки пошуку.",
        f"На {P5['n_valid']} запитах середній P@5: TF-IDF {f(P5['mean']['tfidf'], 2)}, "
        f"BM25 {f(P5['mean']['bm25'], 2)}.",
    ])
    r.h("8. Запуск")
    r.code('''uv sync
uv run python -m findex.index data/gutenberg --out index.pkl --positions
uv run python -m findex.engine index.pkl 'whale AND (sea OR ocean) NOT ship "white whale"' -k 5 -v
uv run python -m findex.lab3_eval data/gutenberg --limit 50
uv run pytest && uv run ruff check .''')
    r.d.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="results_lab3.json")
    ap.add_argument("--out", default="Лабораторна3_звіт.docx")
    a = ap.parse_args()
    with open(a.results, encoding="utf-8") as fh:
        build(json.load(fh), a.out)
    print("→", a.out)


if __name__ == "__main__":
    main()
