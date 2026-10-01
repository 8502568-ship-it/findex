"""Збирає звіт Лаби 5 (docx) і README_lab5.md з results_lab5.json.

    # після вимірювань (усі числа беруться з results_lab5.json):
    uv run python scripts/make_report_lab5.py --results results_lab5.json \
        --out "Лабораторна5_звіт.docx" --readme README_lab5.md
    # чернетка без чисел (замість них підсвічені [X]):
    uv run python scripts/make_report_lab5.py --draft --out "Лабораторна5_чернетка.docx"

Текст — чернетка пояснень: перечитай і переформулюй своїми словами, а висновки
звір зі своїми цифрами (лаба вимагає саме твоїх слів).
"""
from __future__ import annotations

import argparse
import dis
import io
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt
from lab5_data import BASE_LABEL, load, rows

REPO = "https://github.com/8502568-ship-it/findex"
DRAFT = False
Block = tuple[Any, ...]


def n(x: float, d: int = 2) -> str:
    return "[X]" if DRAFT else f"{x:.{d}f}".replace(".", ",")


def x_(x: float) -> str:
    return "[X]" if DRAFT else f"{x:.2f}".replace(".", ",") + "×"


def draft_results() -> dict[str, Any]:
    """Фіктивні дані: лише щоб форма звіту була повною; числа підміняються на [X]."""
    def run(ex: str, w: int, wall: float, cpu: float, merge: float, rss: float, gil: bool) -> dict[str, Any]:
        return {"executor": ex, "workers": w, "wall": wall, "cpu": cpu, "merge": merge,
                "peak_rss_mib": rss, "postings": 1, "docs": 1, "terms": 1, "python": "3.13",
                "gil_enabled": gil}
    ru = {BASE_LABEL: [run("serial", 1, 100, 100, 1, 500, True)], "cpython-3.13t": []}
    for w in (1, 2, 4, 8):
        ru[BASE_LABEL].append(run("threads", w, 105, 105, 1, 520, True))
        ru[BASE_LABEL].append(run("processes", w, 100 / w * 1.2 + 8, 130, 8, 600, True))
        ru["cpython-3.13t"].append(run("threads", w, 110 / w + 8, 120, 8, 600, False))
    meta = {k: {"cpu_model": "[X]", "cpu_count": 8, "platform": "[X]", "python": "[X]",
                "gil_enabled": k == BASE_LABEL, "docs": "[X]", "terms": "[X]"} for k in ru}
    return {"meta": meta, "runs": ru}


# ============================ побудова змісту ===============================
def facts(data: dict[str, Any]) -> dict[str, Any]:
    rs = rows(data)
    base = BASE_LABEL if BASE_LABEL in data["runs"] else next(iter(data["runs"]))
    ft = next((lb for lb in data["runs"] if lb != base), None)
    cores = max(r["workers"] for r in rs if r["label"] == base)
    get = lambda lb, ex, w: next((r for r in rs if (r["label"], r["executor"], r["workers"]) == (lb, ex, w)), None)
    serial = get(base, "serial", 1)
    thr = [r for r in rs if r["label"] == base and r["executor"] == "threads"]
    prc = [r for r in rs if r["label"] == base and r["executor"] == "processes"]
    best_t, best_p = max(thr, key=lambda r: r["speedup"]), max(prc, key=lambda r: r["speedup"])
    pc = get(base, "processes", cores)
    f = pc["merge"] / serial["wall"]
    ftr = [r for r in rs if r["label"] == ft and r["executor"] == "threads"] if ft else []
    return {"rs": rs, "base": base, "ft": ft, "cores": cores, "serial": serial, "thr": thr,
            "prc": prc, "best_t": best_t, "best_p": best_p, "pc": pc, "f": f,
            "amdahl": lambda w: 1 / (f + (1 - f) / w),
            "ft_best": max(ftr, key=lambda r: r["speedup"]) if ftr else None,
            "meta": data["meta"][base], "meta_ft": data["meta"].get(ft) if ft else None}


def read_demo(path: Path) -> str | None:
    return path.read_text(encoding="utf-8").strip() if path.exists() else None


def inc_dis() -> str:
    def inc() -> None:
        global counter
        counter += 1

    buf = io.StringIO()
    dis.dis(inc, file=buf)
    return f"# Python {sys.version.split()[0]}\n" + buf.getvalue().strip()


def table_blocks(F: dict[str, Any]) -> list[Block]:
    body = []
    for r in F["rs"]:
        body.append([r["label"], r["executor"], str(r["workers"]), n(r["wall"]), n(r["cpu"]),
                     n(r["rss"], 0), n(r["merge"]), x_(r["speedup"]) if r["executor"] != "serial" else "1,00×"])
    am = [[str(r["workers"]), x_(r["speedup"]), x_(F["amdahl"](r["workers"]))] for r in F["prc"]]
    return [
        ("cap", "Таблиця 1 — збірка повного корпусу (медіана ≥ 3 запусків; прискорення відносно serial на "
                f"{F['base']})"),
        ("table", ["Інтерпретатор", "Executor", "Workers", "Wall, с", "CPU, с", "Пік RSS батька, МіБ",
                   "Merge, с", "Прискорення"], body),
        ("cap", f"Таблиця 2 — закон Амдала для processes: послідовна частка f = merge / serial wall = "
                f"{n(F['f'] * 100, 1)} %, стеля 1/f = {x_(1 / F['f'])}"),
        ("table", ["Workers", "Виміряно", "Амдал 1/(f+(1−f)/N)"], am),
    ]


def machine_block(F: dict[str, Any]) -> Block:
    m = F["meta"]
    return ("p", f"Машина: {m['cpu_model']}, логічних ядер (os.cpu_count): {m['cpu_count']}, ОС: "
                 f"{m['platform']}. Корпус: {m['docs']} документів, словник {m['terms']} термінів. "
                 f"Python {m['python']} (sys._is_gil_enabled() = {m['gil_enabled']})"
            + (f"; вільнопотокова збірка: Python {F['meta_ft']['python']} "
               f"(sys._is_gil_enabled() = {F['meta_ft']['gil_enabled']})." if F["meta_ft"] else "."))


def method_blocks() -> list[Block]:
    return [
        ("p", "Кожна клітинка (executor × workers) міряється щонайменше тричі; у таблицю йде медіана. "
              "Кожен запуск — окремий свіжий процес Python (python -m findex.bench5 --one), тому пік пам'яті і "
              "CPU-час не успадковують нічого від попередніх запусків. Перед серією робиться один прогрівальний "
              "запуск (холодний кеш файлів не потрапляє в таблицю). Збірка міряється без запису індексу на диск."),
        ("bullets", [
            "Wall — time.perf_counter() від старту збірки до готового індексу (включно з plan_chunks, "
            "передачею результатів воркерів і merge).",
            "CPU — time.process_time() батьківського процесу; для processes додається сума CPU-часу воркерів "
            "(time.thread_time() усередині build_partial). Старт інтерпретатора воркера і імпорти в цю суму не "
            "входять, тож вона трохи занижена.",
            "Пік RSS — найвищий resident set саме батьківського процесу (Windows: PeakWorkingSetSize, Linux: "
            "ru_maxrss). Пам'ять воркерів сюди не входить.",
            "Merge — окремий таймер навколо merge(); це послідовна частина збірки.",
            "Прискорення — wall serial на звичайній збірці CPython поділити на wall клітинки; одна й та сама "
            "база для рядків 3.13t, щоб їх можна було порівнювати.",
            "Шматків: 1 для serial і для workers=1, інакше workers×4 (більше шматків, ніж воркерів, вирівнює "
            "навантаження: файли корпусу дуже різного розміру). Шматки неперервні й зважені за розміром файлів.",
        ]),
    ]


def gil_page(F: dict[str, Any], race: str | None, race_ft: str | None) -> list[Block]:
    bt, bp, pc, s, f = F["best_t"], F["best_p"], F["pc"], F["serial"], F["f"]
    B: list[Block] = []
    thr_flat = bt["speedup"] < 1.15
    B.append(("p", f"Потоки під GIL. Найкраще прискорення потоків — {x_(bt['speedup'])} при {bt['workers']} воркерах. "
                   + ("Тобто практично нуль. Індексація — це токенізація і оновлення словників, тобто чистий Python-"
                      "байткод, а GIL дозволяє виконувати байткод лише одному потоку одночасно. Вісім потоків не "
                      "працюють паралельно: вони по черзі беруть GIL (кожні 5 мс інтерпретатор просить поточний потік "
                      "його віддати) і чекають одне одного. Додатково з'являються витрати на перемикання і гірший "
                      "кеш процесора, тому потоки можуть вийти навіть трохи повільнішими за serial. Читання файлів "
                      "GIL відпускає, але воно становить малу частку часу порівняно з токенізацією."
                      if thr_flat else
                      "Це більше, ніж очікувалось для чистого Python-коду під GIL: перевір, чи не відпускає GIL "
                      "якась частина build_partial (читання файлів) і яку частку часу вона займає, і поясни це число.")))
    B.append(("p", f"Процеси. Найкраще прискорення — {x_(bp['speedup'])} при {bp['workers']} воркерах. Кожен процес має "
                   "власний інтерпретатор і власний GIL, тож токенізація йде на різних ядрах справді паралельно. "
                   "Але це не безкоштовно. Процес запускається методом spawn: новий інтерпретатор імпортує модулі "
                   "заново. Результат воркера (частковий індекс з сотнями тисяч термінів) серіалізується через pickle, "
                   "іде по каналу і десеріалізується в батьківському процесі. Тому сумарний CPU-час "
                   + ("більший, ніж у " if DRAFT or pc["cpu"] > s["cpu"] else "не більший, ніж у ")
                   + f"serial: {n(pc['cpu'])} с проти {n(s['cpu'])} с на {pc['workers']} воркерах, тоді як wall впав до "
                   f"{n(pc['wall'])} с проти {n(s['wall'])} с. Різниця між CPU і wall — ціна паралелізму: ядра зайняті "
                   f"одночасно, а частина їхньої роботи — накладні витрати. Пам'ять: пік RSS батьківського "
                   f"процесу {n(pc['rss'], 0)} МіБ проти {n(s['rss'], 0)} МіБ у serial (батько тримає отримані часткові "
                   "індекси і збирає з них результат), а кожен воркер додатково тримає власну копію інтерпретатора "
                   "і свій частковий індекс — їх RSS у таблиці не враховано, тож реальне споживання пам'яті системою "
                   "більше за цю колонку."))
    B.append(("p", f"Закон Амдала. Merge виконується в одному процесі, поки решта чекає. На {pc['workers']} воркерах він займає "
                   f"{n(pc['merge'])} с, тобто f = {n(f * 100, 1)} % від wall serial ({n(s['wall'])} с). За законом Амдала "
                   f"S(N) = 1/(f + (1−f)/N), тому навіть з нескінченною кількістю ядер прискорення не перевищить "
                   f"1/f = {x_(1 / f)}; для {pc['workers']} воркерів формула дає {x_(F['amdahl'](pc['workers']))}, "
                   f"а виміряно {x_(pc['speedup'])} (див. таблицю 2). Це лише нижня оцінка послідовної частки: "
                   "в неї не входить десеріалізація результатів воркерів у батьківському процесі і запуск процесів, "
                   "тому реальна стеля нижча; різницю між формулою і виміром варто пояснити саме цим."))
    ftb = F["ft_best"]
    if ftb:
        B.append(("p", f"Збірка 3.13t. На вільнопотоковому інтерпретаторі (sys._is_gil_enabled() = "
                       f"{F['meta_ft']['gil_enabled']}) ті самі ThreadPoolExecutor-потоки дали {x_(ftb['speedup'])} при "
                       f"{ftb['workers']} воркерах (процеси на звичайній збірці — {x_(bp['speedup'])}). Потоки тепер "
                       "справді виконують байткод паралельно, не платять за pickle і працюють у спільній пам'яті, "
                       "тож CPU-час і RSS мають бути ближчими до serial, ніж у processes — звір це з таблицею 1. "
                       "Merge лишається послідовним, тому Амдал діє й тут. Ціна — повільніший одноядерний "
                       "інтерпретатор (дивись рядок serial на 3.13t, якщо він є)."))
    if race:
        B.append(("p", "Гонитва даних. Без GIL спільний лічильник, який оновлюють кілька потоків, губить оновлення. "
                       "Під GIL це могло не помічатися лише тому, що перемикання між байткодами траплялося рідко, "
                       "а не тому що код був коректним. Результати scripts/gil_demo.py:"))
        B.append(("code", f"GIL-збірка:\n{race}" + (f"\n\nВільнопотокова збірка:\n{race_ft}" if race_ft else "")))
        B.append(("p", "Причина: counter += 1 розкладається на кілька байткодів (прочитати, додати, записати), і потік, "
                       "що встиг прочитати старе значення, затирає чужий приріст. Фікс: threading.Lock навколо оновлення "
                       "або, краще, відсутність спільного стану: кожен потік веде власний Counter, а результати зливаються "
                       "в кінці. Саме так працює build_partial: кожен воркер будує власний InvertedIndex, спільних "
                       "змінних немає, тому збірка потоками коректна й без GIL, а merge виконується в одному потоці."))
    return B


def reflection(F: dict[str, Any]) -> list[Block]:
    pc, s = F["pc"], F["serial"]
    q = [
        ("1. Що таке GIL, що він блокує і чого не захищає?",
         "GIL — м'ютекс усередині CPython: щоб виконувати Python-байткод, потік мусить його утримувати, тож "
         "одночасно байткод виконує один потік. Він потрібен, бо пам'ять керується підрахунком посилань, і одночасні "
         "інкременти/декременти лічильника посилань з кількох потоків його б зіпсували. GIL захищає внутрішній стан "
         "інтерпретатора, а не інваріанти програми: counter += 1 лишається неатомарним, складені дії «перевірив — "
         "змінив» потребують Lock."),
        ("2. Що робили 8 потоків, коли збірка не прискорилась?",
         "Один тримав GIL і виконував токенізацію, сімох чекали на GIL. Кожні ~5 мс (sys.getswitchinterval) "
         "поточний потік просять його віддати, і GIL переходить до іншого. Паралельно працювало лише читання файлів, "
         "бо воно відпускає GIL. Корисної роботи за секунду не більшає, а накладні витрати на перемикання і "
         "конкуренцію за GIL додаються."),
        ("3. Куди пішли «зайві» CPU-секунди у процесів?",
         f"Приклад з наших вимірів: {pc['workers']} воркерів дали wall {n(pc['wall'])} с і CPU {n(pc['cpu'])} с (serial: "
         f"{n(s['wall'])} с і {n(s['cpu'])} с). CPU-час — сума по всіх ядрах, тому він і має бути більшим за wall. "
         "Понад serial додаються: запуск воркерів (spawn, повторні імпорти), pickle результатів у воркері, "
         "unpickle в батьку, конкуренція за кеш і пам'ять, а також merge, який виконує лише батько."),
        ("4. Чому функція для ProcessPoolExecutor має бути на рівні модуля? А lambda?",
         "Функцію пересилають у воркер як посилання: pickle зберігає ім'я модуля і ім'я функції, а воркер її "
         "імпортує. Лямбда й вкладена функція не мають імпортованого імені, тому pickle падає (PicklingError або "
         "AttributeError: Can't pickle local object)."),
        ("5. fork проти spawn: що робить кожен, у чому небезпека fork, навіщо __main__-guard?",
         "fork копіює батьківський процес разом з пам'яттю — швидко, але копіюється й стан потоків і блокувань: якщо "
         "в момент fork якийсь інший потік тримав lock, у дитини він залишиться заблокованим назавжди. spawn "
         "запускає чистий інтерпретатор і імпортує головний модуль заново. Без захисту `if __name__ == \"__main__\"` "
         "це повторне імпортування знову запустило б пул, і процеси плодилися б рекурсивно (Python це помічає і "
         "кидає RuntimeError). Ми явно просимо spawn, а в CLI запуск іде через entry point typer, що еквівалентно."),
        ("6. Закон Амдала; merge = 15 % — яка стеля?",
         "S(N) = 1 / (f + (1 − f)/N), де f — послідовна частка. При N → ∞ маємо S = 1/f. Для f = 0,15 це "
         f"1/0,15 ≈ 6,67×. У наших вимірах f ≈ {n(F['f'] * 100, 1)} % (merge/wall serial), стеля {x_(1 / F['f'])}."),
        ("7. Чому вільнопотокова збірка виявила баг, якого не видно на звичайній? Чи коректний код на звичайній?",
         "GIL випадково серіалізував багато операцій, і неатомарний read-modify-write рідко переривався посередині. "
         "Без GIL потоки справді виконуються одночасно, і втрати оновлень стають частими. Код на звичайній збірці "
         "не коректний, а лише щасливий: його коректність залежить від деталей планувальника і версії інтерпретатора, "
         "а не від гарантій мови."),
        ("8. dis для counter += 1 і місце втрати оновлення",
         "Байткод (нижче): значення читається (LOAD_GLOBAL), додається (BINARY_OP) і записується (STORE_GLOBAL). Якщо "
         "потік A прочитав 5 і його перервали, потік B читає 5, пише 6, потім A дописує свої 6 — один приріст втрачено."),
        ("9. Де потоки, де процеси, де ні те, ні інше?",
         "Потоки: очікування вводу-виводу (завантаження десятків URL, читання файлів), бо GIL під час очікування "
         "відпущено. Процеси: CPU-bound чистий Python — наша індексація. Ні те, ні інше: тисячі одночасних "
         "з'єднань, де потік на кожне надто дорогий — там asyncio (Лаба 6)."),
    ]
    out: list[Block] = []
    for title, ans in q:
        out += [("h3", title), ("p", ans)]
        if title.startswith("8."):
            out.append(("code", inc_dis()))
    return out


def readme_blocks(F: dict[str, Any], race: str | None, race_ft: str | None, img: str) -> list[Block]:
    return [
        ("h1", "Лабораторна 5 — конкурентність і GIL: паралельна індексація"),
        ("h2", "Запуск"),
        ("code", "uv sync\nuv run findex index data/gutenberg --workers 4 --executor processes -o index.json\n"
                 "uv run python -m findex.bench5 data/gutenberg --repeats 3          # звичайний CPython\n"
                 "uv run python scripts/gil_demo.py                                  # GIL, гонитва, I/O\n"
                 "uv run python scripts/plot_lab5.py\n"
                 "uv run pytest"),
        ("h2", "Машина і методика"), machine_block(F), *method_blocks(),
        ("h2", "Результати"), *table_blocks(F),
        ("img", img, "Рисунок 1 — прискорення збірки індексу від числа воркерів (пунктир — ідеальне лінійне)"),
        ("h2", "Чому так: GIL, процеси, Амдал, 3.13t"), *gil_page(F, race, race_ft),
    ]


def report_blocks(F: dict[str, Any], race: str | None, race_ft: str | None, img: str) -> list[Block]:
    return [
        ("title", "Лабораторна робота 5. Конкурентність і GIL: потоки, процеси і паралельна індексація"),
        ("p", f"Проєкт: findex. Репозиторій: {REPO}, тег lab-05."),
        ("h2", "1. Що реалізовано"),
        ("bullets", [
            "build_index розділено на build_partial(paths, first_id, positions) і merge(partials) "
            "(src/findex/parallel.py). build_partial — функція рівня модуля; у воркер ідуть лише рядки зі шляхами і числа.",
            "doc_id унікальні на весь корпус: plan_chunks наперед ріже відсортований список файлів на неперервні "
            "шматки і видає кожному first_id; нумерація збігається з послідовною (enumerate) із Лаби 4.",
            "Один і той самий код працює під простим циклом, ThreadPoolExecutor і ProcessPoolExecutor (spawn): "
            "findex index --workers N --executor {serial,threads,processes}.",
            "Падіння воркера валить збірку: перший виняток піднімається з future.result(), решта черги скасовується, "
            "для processes у винятку лишається стек воркера; CLI друкує його і виходить з кодом 1.",
            "Тести (tests/test_lab5.py): merge([build_partial(усі)]) дорівнює індексу Лаби 4, включно з порядком "
            "ключів; усі три executor дають ідентичний індекс; CLI-файли однакові байт у байт; падіння воркера; "
            "битий UTF-8 пропускається однаково у всіх режимах.",
            "Виправлено помилку в репозиторії: index.py імпортував неіснуючий модуль findex.tokenizer, тепер "
            "findex.tokenize.",
        ]),
        ("h2", "2. Машина і методика"), machine_block(F), *method_blocks(),
        ("h2", "3. Результати"), *table_blocks(F),
        ("img", img, "Рисунок 1 — прискорення від числа воркерів"),
        ("h2", "4. Пояснення: GIL, процеси, Амдал, 3.13t"), *gil_page(F, race, race_ft),
        ("h2", "5. Висновок щодо executor за замовчуванням"),
        ("p", (f"Процеси виграли ({x_(F['best_p']['speedup'])} при {F['best_p']['workers']} воркерах), тому типовим "
               "executor стає processes (константа DEFAULT_EXECUTOR у parallel.py).")
         if DRAFT or F["best_p"]["speedup"] > 1.0 else
         "Процеси на цій машині не виграли, тому типовим лишається serial."),
        ("h2", "6. Відповіді на Reflection"), *reflection(F),
    ]


# ================================ рендер ====================================
def to_md(blocks: list[Block]) -> str:
    out: list[str] = []
    for b in blocks:
        k = b[0]
        if k in ("h1", "title"):
            out.append(f"# {b[1]}\n")
        elif k == "h2":
            out.append(f"## {b[1]}\n")
        elif k == "h3":
            out.append(f"### {b[1]}\n")
        elif k == "p":
            out.append(f"{b[1]}\n")
        elif k == "cap":
            out.append(f"**{b[1]}**\n")
        elif k == "bullets":
            out.append("\n".join(f"- {t}" for t in b[1]) + "\n")
        elif k == "code":
            out.append(f"```\n{b[1]}\n```\n")
        elif k == "table":
            head, body = b[1], b[2]
            out.append("| " + " | ".join(head) + " |\n|" + "---|" * len(head))
            out += ["| " + " | ".join(r) + " |" for r in body]
            out.append("")
        elif k == "img":
            out.append(f"![{b[2]}]({b[1]})\n\n*{b[2]}*\n")
    return "\n".join(out)


def _add_runs(par: Any, text: str, size: int = 12, bold: bool = False, mono: bool = False) -> None:
    for i, piece in enumerate(text.split("[X]")):
        for j, t in enumerate((piece,) if i == 0 else ("[X]", piece)):
            if not t:
                continue
            run = par.add_run(t)
            run.font.size = Pt(size)
            run.font.name = "Consolas" if mono else "Times New Roman"
            run.bold = bold
            if t == "[X]" and j == 0 and i > 0:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def to_docx(blocks: list[Block], out: Path, img_path: Path | None) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    for b in blocks:
        k = b[0]
        if k == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(p, b[1], 16, bold=True)
        elif k in ("h1", "h2", "h3"):
            _add_runs(doc.add_paragraph(), b[1], {"h1": 16, "h2": 14, "h3": 12}[k], bold=True)
        elif k == "p":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(1.25)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(p, b[1])
        elif k == "cap":
            p = doc.add_paragraph()
            _add_runs(p, b[1], 11, bold=True)
        elif k == "bullets":
            for t in b[1]:
                _add_runs(doc.add_paragraph(style="List Bullet"), t)
        elif k == "code":
            for line in b[1].splitlines() or [""]:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                _add_runs(p, line, 9, mono=True)
            doc.add_paragraph()
        elif k == "table":
            t = doc.add_table(rows=1, cols=len(b[1]), style="Table Grid")
            for c, h in zip(t.rows[0].cells, b[1]):
                c.text = ""
                _add_runs(c.paragraphs[0], h, 10, bold=True)
            for row in b[2]:
                cells = t.add_row().cells
                for c, v in zip(cells, row):
                    c.text = ""
                    _add_runs(c.paragraphs[0], v, 10)
            doc.add_paragraph()
        elif k == "img":
            if img_path and img_path.exists():
                doc.add_picture(str(img_path), width=Cm(14))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                _add_runs(doc.add_paragraph(), "[X] — тут має бути графік (запусти scripts/plot_lab5.py)")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(p, b[2], 11)
    doc.save(out)


def main() -> None:
    global DRAFT
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--results", type=Path, default=Path("results_lab5.json"))
    ap.add_argument("--out", type=Path, default=Path("Лабораторна5_звіт.docx"))
    ap.add_argument("--readme", type=Path, default=None, help="записати ще й README_lab5.md")
    ap.add_argument("--image", type=Path, default=Path("docs/speedup_lab5.png"))
    ap.add_argument("--draft", action="store_true", help="числа замінити на підсвічені [X]")
    args = ap.parse_args()
    DRAFT = args.draft
    data = draft_results() if DRAFT else load(args.results)
    F = facts(data)
    race = read_demo(Path("docs/gil_demo_gil.txt"))
    race_ft = read_demo(Path("docs/gil_demo_free.txt"))
    img_rel = args.image.as_posix()
    to_docx(report_blocks(F, race, race_ft, img_rel), args.out, args.image)
    print(f"saved -> {args.out}")
    if args.readme:
        args.readme.write_text(to_md(readme_blocks(F, race, race_ft, img_rel)), encoding="utf-8")
        print(f"saved -> {args.readme}")


if __name__ == "__main__":
    main()
